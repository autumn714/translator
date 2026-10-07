"""Per-format round trips through the JobManager with fake translators."""
from __future__ import annotations

import io
import re
import zipfile

import docfixtures as F
import pytest

from translator_app.documents.base import (
    MSG_HWP_AS_DOCX,
    MSG_LLM_UNAVAILABLE,
    MSG_NO_TEXT,
    MSG_VISION_UNAVAILABLE,
    DocumentError,
)
from translator_app.documents.jobs import JobManager
from translator_app.documents.ooxml import DOCX_PARTS


@pytest.fixture
def anyio_backend():
    return "asyncio"


@pytest.fixture
async def make_manager(tmp_path):
    managers: list[JobManager] = []

    async def factory(translator=None, **over):
        m = JobManager(settings=F.settings(tmp_path / f"m{len(managers)}", **over),
                       translator=translator or F.FakeTranslator())
        await m.start()
        managers.append(m)
        return m

    yield factory
    for m in managers:
        await m.aclose()


def _untouched_members_equal(src: bytes, out: bytes, changed: re.Pattern) -> None:
    with zipfile.ZipFile(io.BytesIO(src)) as a, zipfile.ZipFile(io.BytesIO(out)) as b:
        ia, ib = a.infolist(), b.infolist()
        assert [i.filename for i in ia] == [i.filename for i in ib]
        for x, y in zip(ia, ib):
            assert x.compress_type == y.compress_type, x.filename
            if not changed.fullmatch(x.filename):
                assert (x.CRC, x.file_size) == (y.CRC, y.file_size), x.filename


# ====================================================================== DOCX
@pytest.mark.anyio
async def test_docx_translated(make_manager):
    import docx

    tr = F.FakeTranslator()
    m = await make_manager(tr)
    src = F.make_docx()
    job, out = await F.run_job(m, "매뉴얼.docx", src)
    assert job.status == "done", job.error
    assert job.output_filename == "매뉴얼_ko.docx"
    xml = F.zip_text(out, "word/document.xml")
    assert "[KO] Electrolyzer operation" in xml
    assert " PAGE " in xml and xml.count("instrText") == 2      # field code untouched
    assert ">12345<" in xml and ">700<" in xml                   # numbers are not sent
    headers = [n for n in F.zip_names(out) if re.fullmatch(r"word/header\d*\.xml", n)]
    assert any("[KO] Hydrogen safety manual" in F.zip_text(out, n) for n in headers)
    _untouched_members_equal(src, out, re.compile(DOCX_PARTS.pattern))

    d = docx.Document(io.BytesIO(out))
    para = next(p for p in d.paragraphs if "electrolyzer" in p.text)
    assert para.text.startswith("[KO] The ")
    assert any(r.bold and r.text == "electrolyzer" for r in para.runs)
    assert any(r.italic and r.text == "30 bar" for r in para.runs)
    assert d.tables[0].cell(1, 0).text == "[KO] Storage pressure"
    # tagged strings were sent with tags=True and document priority
    assert all(c["priority"] == "document" for c in tr.calls)
    assert any(c["tags"] and any("<g1>" in t for t in c["texts"]) for c in tr.calls)
    assert job.report_count == 0


@pytest.mark.anyio
@pytest.mark.parametrize("mode", ["reorder", "drop_tags"])
async def test_docx_tag_variants(make_manager, mode):
    import docx

    tr = F.FakeTranslator(mode)
    m = await make_manager(tr)
    job, out = await F.run_job(m, "a.docx", F.make_docx())
    assert job.status == "done", job.error
    d = docx.Document(io.BytesIO(out))
    para = next(p for p in d.paragraphs if "electrolyzer" in p.text)
    assert para.text.startswith("[KO]")
    items = m.report_items(job)
    if mode == "reorder":
        assert any(r.bold and r.text == "electrolyzer" for r in para.runs)
        assert not any("서식" in i["issue"] for i in items)
    else:
        assert "1,250.5" in para.text
        assert any("서식" in i["issue"] for i in items)


@pytest.mark.anyio
async def test_docx_bilingual(make_manager):
    import docx

    m = await make_manager()
    job, out = await F.run_job(m, "a.docx", F.make_docx(), output="bilingual")
    assert job.status == "done", job.error
    assert job.output_filename == "a_en-ko.docx"
    xml = F.zip_text(out, "word/document.xml")
    assert xml.count("instrText") == 2                            # the field is not copied
    texts = [p.text for p in docx.Document(io.BytesIO(out)).paragraphs]
    i = texts.index("Electrolyzer operation")
    assert texts[i + 1] == "[KO] Electrolyzer operation"
    j = next(k for k, t in enumerate(texts) if t.startswith("The electrolyzer"))
    assert texts[j + 1].startswith("[KO] The electrolyzer")
    cell = docx.Document(io.BytesIO(out)).tables[0].cell(1, 0)
    assert [p.text for p in cell.paragraphs] == ["Storage pressure", "[KO] Storage pressure"]


@pytest.mark.anyio
async def test_selftest_minimal_packages(make_manager):
    """translator_app.selftest uploads DOCX/XLSX without styles.xml/docProps/calcPr."""
    m = await make_manager()
    src = F.make_minimal_docx(["Hydrogen safety notes", "Check the pressure gauge before every use."])
    job, out = await F.run_job(m, "selftest.docx", src)
    assert job.status == "done", job.error
    assert F.zip_text(out, "word/document.xml") != F.zip_text(src, "word/document.xml")
    assert "[KO] Hydrogen safety notes" in F.zip_text(out, "word/document.xml")

    xsrc = F.make_minimal_xlsx()
    job, out = await F.run_job(m, "selftest.xlsx", xsrc)
    assert job.status == "done", job.error
    assert "[KO] Hydrogen storage tank" in F.zip_text(out, "xl/sharedStrings.xml")
    assert 'fullCalcOnLoad="1"' in F.zip_text(out, "xl/workbook.xml")


# ====================================================================== PPTX / XLSX
@pytest.mark.anyio
async def test_pptx(make_manager):
    from pptx import Presentation

    m = await make_manager()
    src = F.make_pptx()
    job, out = await F.run_job(m, "deck.pptx", src)
    assert job.status == "done", job.error
    prs = Presentation(io.BytesIO(out))
    slide = prs.slides[0]
    texts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]
    assert "[KO] Hydrogen roadmap 2030" in texts
    body = next(t for t in texts if "Capacity" in t)
    assert "[KO] Capacity will reach 12 GW" in body and "always" in body
    table = next(sh.table for sh in slide.shapes if sh.has_table)
    assert table.cell(1, 0).text == "[KO] Europe" and table.cell(1, 1).text == "35%"
    assert slide.notes_slide.notes_text_frame.text == "[KO] Speaker note about storage"
    assert 'lang="ko-KR"' in F.zip_text(out, "ppt/slides/slide1.xml")
    bold = [r for p in slide.placeholders[1].text_frame.paragraphs for r in p.runs if r.font.bold]
    assert [r.text for r in bold] == ["Safety "]
    _untouched_members_equal(src, out, re.compile(r"ppt/(slides/slide|notesSlides/notesSlide)\d+\.xml"))


@pytest.mark.anyio
async def test_xlsx_formulas_and_literals(make_manager):
    from openpyxl import load_workbook

    m = await make_manager()
    job, out = await F.run_job(m, "plants.xlsx", F.make_xlsx())
    assert job.status == "done", job.error
    wb = load_workbook(io.BytesIO(out))
    ws = wb["Plants"]                                            # sheet names untouched
    assert ws["A2"].value == "[KO] North plant"
    assert ws["B2"].value == "Yes"                               # used as a formula literal
    assert ws["D2"].value == '=IF(B2="Yes","Running","Stopped")'
    assert ws["C2"].value == 120
    assert ws["A2"].comment.text == "[KO] Commissioned in 2024"
    assert 'fullCalcOnLoad="1"' in F.zip_text(out, "xl/workbook.xml")


# ====================================================================== HWPX / HWP
@pytest.mark.anyio
async def test_hwpx(make_manager):
    m = await make_manager()
    src = F.make_hwpx()
    job, out = await F.run_job(m, "보고서.hwpx", src)
    assert job.status == "done", job.error
    infos = zipfile.ZipFile(io.BytesIO(out)).infolist()
    assert infos[0].filename == "mimetype" and infos[0].compress_type == zipfile.ZIP_STORED
    sec = F.zip_text(out, "Contents/section0.xml")
    assert "[KO] The " in sec and ">electrolyzer<" in sec
    assert '<hp:tab width="4000" leader="0" type="1"/>' in sec and "<hp:lineBreak/>" in sec
    assert "[KO] Storage tank" in sec
    # stale layout cache removed where the text changed, kept elsewhere
    from lxml import etree

    root = etree.fromstring(zipfile.ZipFile(io.BytesIO(out)).read("Contents/section0.xml"))
    hp = "{%s}" % F.HP_NS
    paras = {p.get("id"): p for p in root.findall(f"{hp}p")}
    assert paras["1"].find(f"{hp}linesegarray") is None
    assert paras["3"].find(f"{hp}linesegarray") is not None   # "2024" was not translated
    prv = zipfile.ZipFile(io.BytesIO(out)).read("Preview/PrvText.txt").decode("utf-8")
    assert "[KO] The electrolyzer" in prv
    _untouched_members_equal(src, out, re.compile(r"Contents/section\d+\.xml|Preview/PrvText\.txt"))


@pytest.mark.anyio
async def test_hwpx_bilingual(make_manager):
    m = await make_manager()
    job, out = await F.run_job(m, "a.hwpx", F.make_hwpx(), output="bilingual")
    assert job.status == "done", job.error
    sec = F.zip_text(out, "Contents/section0.xml")
    assert sec.index("> produces hydrogen at 30 bar.<") < sec.index("[KO] The ")
    assert sec.count("<hp:tbl ") == 1                             # tables are not duplicated
    assert "Storage tank</hp:t>" in sec and "[KO] Storage tank" in sec
    assert "linesegarray" not in sec


@pytest.mark.anyio
@pytest.mark.parametrize("output", ["translated", "bilingual"])
async def test_hwp_to_docx(make_manager, output):
    import docx

    sample = F.FIXTURES / "sample.hwp"
    if not sample.exists():
        pytest.skip("HWP fixture missing")
    m = await make_manager()
    job, out = await F.run_job(m, "구버전.hwp", sample.read_bytes(), output=output)
    assert job.status == "done", job.error
    assert job.output_filename.endswith(".docx")
    assert MSG_HWP_AS_DOCX in job.warnings
    texts = [p.text for p in docx.Document(io.BytesIO(out)).paragraphs]
    assert "[KO] Clean Hydrogen Certification Guide" in texts
    if output == "bilingual":
        i = texts.index("Clean Hydrogen Certification Guide")
        assert texts[i + 1] == "[KO] Clean Hydrogen Certification Guide"
    assert "[KO] Clean Hydrogen" in m.preview_text(job)


@pytest.mark.anyio
async def test_hwp_rejects_non_hwp(make_manager):
    m = await make_manager()
    job, _ = await F.run_job(m, "fake.hwp", F.make_hwpx())
    assert job.status == "error" and ".hwpx" in job.error


# ====================================================================== PDF
@pytest.mark.anyio
async def test_pdf_layout_with_scanned_pages(make_manager):
    import pymupdf

    tr = F.FakeTranslator()
    m = await make_manager(tr)
    src = F.make_pdf(scanned=True, ocr_layer=True)
    job, out = await F.run_job(m, "report.pdf", src)
    assert job.status == "done", job.error
    pages = F.pdf_text(out)
    assert len(pages) == 5                                        # 3 pages + 2 inserted translations
    assert "[KO] Clean Hydrogen Market Outlook" in pages[0]
    assert "Global electrolyzer capacity reached" not in pages[0].replace("[KO] Global electrolyzer capacity", "")
    assert pages[1].strip() == ""                                  # scanned page untouched
    assert pages[2].startswith("p.2 번역") and "번역된 제목" in pages[2]
    assert "[KO]" not in pages[3]                                  # OCR-layer page untouched
    assert pages[4].startswith("p.3 번역")
    with pymupdf.open(stream=src, filetype="pdf") as a, pymupdf.open(stream=out, filetype="pdf") as b:
        assert len(a[0].get_images()) == len(b[0].get_images())   # images kept
        assert len(a[0].get_drawings()) == len(b[0].get_drawings())
    assert [c["mime"] for c in tr.image_calls] == ["image/png", "image/png"]
    assert any("스캔 페이지" in w for w in job.warnings)
    assert job.total == 2 + sum(len(call["texts"]) for call in tr.calls)   # text units + 2 vision pages
    assert job.public().progress.percent == 100.0
    assert "번역된 제목" in m.preview_text(job)


@pytest.mark.anyio
async def test_pdf_optional_font_dir(make_manager):
    """An admin-provided {data_dir}/fonts/NanumGothic.ttf is used through a pymupdf.Archive."""
    import pymupdf

    from translator_app.documents.pdf import load_font

    m = await make_manager()
    m.font_dir.mkdir(parents=True, exist_ok=True)
    (m.font_dir / "NanumGothic.ttf").write_bytes(pymupdf.Font("cjk").buffer)   # any Hangul TTF will do
    assert load_font(m.font_dir).family == "tfont"
    assert load_font(None).family == "sans-serif"
    job, out = await F.run_job(m, "a.pdf", F.make_pdf(scanned=False))
    assert job.status == "done", job.error
    assert "[KO] Clean Hydrogen Market Outlook" in F.pdf_text(out)[0]


@pytest.mark.anyio
async def test_pdf_docx_mode(make_manager):
    import docx

    m = await make_manager()
    job, out = await F.run_job(m, "report.pdf", F.make_pdf(scanned=True), pdf_mode="docx")
    assert job.status == "done", job.error
    assert job.output_filename == "report_ko.docx"
    d = docx.Document(io.BytesIO(out))
    texts = [p.text for p in d.paragraphs]
    assert "[KO] Clean Hydrogen Market Outlook" in texts
    assert "번역된 제목" in texts and "p.2 번역" in texts
    assert d.tables and d.tables[0].cell(1, 1).text == "30 bar"


@pytest.mark.anyio
async def test_pdf_without_vision(make_manager):
    m = await make_manager(F.FakeTranslator(vision=False))
    job, out = await F.run_job(m, "report.pdf", F.make_pdf(scanned=True, ocr_layer=True))
    assert job.status == "done", job.error
    assert any(MSG_VISION_UNAVAILABLE in w and "2쪽" in w for w in job.warnings)
    pages = F.pdf_text(out)
    # the OCR-layer page is translated from its invisible text layer instead
    assert len(pages) == 4 and pages[3].startswith("p.3 번역") and "[KO] Clean Hydrogen" in pages[3]


@pytest.mark.anyio
async def test_pdf_encrypted_and_broken(make_manager):
    import pymupdf

    doc = pymupdf.open(stream=F.make_pdf(scanned=False), filetype="pdf")
    enc = doc.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="secret", owner_pw="owner")
    doc.close()
    m = await make_manager()
    job, _ = await F.run_job(m, "locked.pdf", enc)
    assert job.status == "error" and "암호" in job.error
    job, _ = await F.run_job(m, "broken.pdf", b"%PDF-1.7 not really")
    assert job.status == "error" and job.error


# ====================================================================== images
@pytest.mark.anyio
@pytest.mark.parametrize("ext", [".png", ".jpg", ".webp"])
async def test_image_to_docx(make_manager, ext):
    import docx
    from PIL import Image

    buf = io.BytesIO()
    Image.open(io.BytesIO(F.make_png())).convert("RGB").save(buf, {"png": "PNG", "jpg": "JPEG", "webp": "WEBP"}[ext[1:]])
    tr = F.FakeTranslator()
    m = await make_manager(tr)
    job, out = await F.run_job(m, f"photo{ext}", buf.getvalue())
    assert job.status == "done", job.error
    assert job.output_filename == "photo_ko.docx"
    d = docx.Document(io.BytesIO(out))
    assert "번역된 제목" in [p.text for p in d.paragraphs]
    assert any(n.startswith("word/media/") for n in F.zip_names(out))
    assert m.preview_text(job).startswith("# 번역된 제목")
    assert tr.image_calls[0]["mime"] == {".png": "image/png", ".jpg": "image/jpeg", ".webp": "image/webp"}[ext]


@pytest.mark.anyio
async def test_image_without_vision_is_an_error(make_manager):
    m = await make_manager(F.FakeTranslator(vision=False))
    job, out = await F.run_job(m, "photo.png", F.make_png())
    assert job.status == "error" and job.error == MSG_VISION_UNAVAILABLE and out is None


# ====================================================================== simple formats
@pytest.mark.anyio
async def test_html(make_manager):
    m = await make_manager()
    job, out = await F.run_job(m, "page.html", F.HTML_DOC)
    assert job.status == "done", job.error
    html = out.decode("utf-8")
    assert html.startswith("<!DOCTYPE html>") and '<html lang="ko">' in html
    assert "<title>[KO] Safety page</title>" in html and 'content="[KO] Hydrogen safety rules"' in html
    assert "<b>valve</b>" in html and 'alt="[KO] Valve photo"' in html and '<a href="tel:119">' in html
    assert "<pre>do_not_translate()</pre>" in html and ">KEEI<" in html


@pytest.mark.anyio
async def test_markdown_bilingual(make_manager):
    m = await make_manager()
    job, out = await F.run_job(m, "notes.md", F.MD_DOC.encode(), output="bilingual")
    assert job.status == "done", job.error
    md = out.decode("utf-8")
    assert md.startswith("---\ntitle: test\n---\n")
    assert "# Safety rules\n\n# [KO] Safety rules" in md
    assert "`valve_1`" in md and "(https://example.com/m.pdf)" in md
    assert "```\ncode block\n```" in md
    assert "| Item<br>[KO] Item | Value<br>[KO] Value |\n|---|---|" in md


@pytest.mark.anyio
async def test_txt_srt_vtt(make_manager):
    m = await make_manager()
    job, out = await F.run_job(m, "notes.txt", F.TXT_DOC.encode("cp949"))
    assert job.status == "done", job.error
    txt = out.decode("utf-8-sig")
    assert txt == "[KO] Safety notes\r\n\r\n  [KO] Indented line about hydrogen\r\n12345\r\n[KO] Last line"
    job, out = await F.run_job(m, "notes.txt", F.TXT_DOC.encode(), output="bilingual")
    assert out.decode().startswith("Safety notes\r\n[KO] Safety notes\r\n")

    job, out = await F.run_job(m, "a.srt", F.SRT_DOC.encode())
    srt = out.decode("utf-8")
    assert "00:00:01,000 --> 00:00:03,000\r\n[KO] Hello <i>world</i>\r\n" in srt and srt.startswith("1\r\n")
    job, out = await F.run_job(m, "a.vtt", F.VTT_DOC.encode())
    vtt = out.decode("utf-8")
    assert vtt.startswith("WEBVTT\n\nNOTE comment stays\n") and "align:start\n[KO] Welcome to the plant" in vtt


# ====================================================================== errors
@pytest.mark.anyio
async def test_unsupported_and_legacy_formats(make_manager):
    m = await make_manager()
    with pytest.raises(DocumentError) as e:
        await m.submit(b"\xd0\xcf\x11\xe0", "old.doc", {})
    assert e.value.status_code == 415 and ".docx" in e.value.message
    with pytest.raises(DocumentError) as e:
        await m.submit(b"x", "virus.exe", {})
    assert e.value.status_code == 415


@pytest.mark.anyio
async def test_protected_and_mislabelled_files(make_manager):
    m = await make_manager()
    ole = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 1024            # encrypted OOXML is an OLE container
    job, _ = await F.run_job(m, "secret.docx", ole)
    assert job.status == "error" and "암호" in job.error
    job, _ = await F.run_job(m, "wrong.docx", F.make_xlsx())
    assert job.status == "error" and "형식" in job.error
    job, _ = await F.run_job(m, "numbers.docx", F.make_minimal_docx(["123", "4.5 %"], sentence=False))
    assert job.status == "error" and job.error == MSG_NO_TEXT


@pytest.mark.anyio
async def test_model_server_down(make_manager):
    m = await make_manager(F.FakeTranslator("unavailable"))
    job, out = await F.run_job(m, "a.docx", F.make_docx())
    assert job.status == "error" and job.error == MSG_LLM_UNAVAILABLE and out is None
    assert not list(m.job_dir(job.id).glob("output*"))
