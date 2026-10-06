"""Regression tests for the document-format fixes (encodings, bilingual copies,
XLSX tables / memory, report rules, Markdown / subtitle edge cases)."""
from __future__ import annotations

import io
import zipfile

import docfixtures as F
import pytest
from lxml import etree

from translator_app.documents import ooxml
from translator_app.documents.base import Collector, preview_from_pairs
from translator_app.documents.html import translate_html
from translator_app.documents.hwpx import HH, translate_hwpx
from translator_app.documents.inline_tags import translate_paragraphs
from translator_app.documents.ooxml import DocxAdapter, docx_clone_paragraph, translate_pptx, translate_xlsx, w
from translator_app.documents.plaintext import decode_text, translate_markdown, translate_subtitles
from translator_app.documents.report import build_report, missing_numbers


def _prefix(xs):
    return ["[KO] " + x for x in xs]


def _collect(fn, *args, **kw):
    c = Collector()
    fn(*args, c, **kw)
    return c.items


# ====================================================================== encodings
@pytest.mark.parametrize(("text", "enc"), [
    ("안녕하세요. 이 문서는 청정수소 인증 제도에 관한 안내서입니다.", "cp949"),
    ("这是一个中文文本文件，用于测试编码检测。我们的产品在国内市场上有很大的发展。", "gbk"),
    ("這是一個中文文本文件，用於測試編碼檢測。我們的產品在國內市場上有很大的發展。", "big5"),
    ("これは日本語のテキストファイルです。文字コードの判定を確認します。", "shift_jis"),
    ("これは日本語のテキストファイルです。文字コードの判定を確認します。", "euc_jp"),
    ("Le café est très chaud. Größe über alles, à bientôt.", "cp1252"),
])
def test_legacy_text_encodings(text, enc):
    decoded, bom = decode_text(text.encode(enc))
    assert decoded == text and bom == "\ufeff"


def test_html_utf8_without_meta_charset():
    src = "<p>안녕하세요 세계</p>".encode()
    assert _collect(translate_html, src) == ["안녕하세요 세계"]
    out = translate_html(src, _prefix).decode("utf-8")
    assert "<p>[KO] 안녕하세요 세계</p>" in out
    assert '<meta charset="utf-8">' in out                       # browsers must not guess Latin-1


@pytest.mark.parametrize("head", [
    '<meta http-equiv="Content-Type" content="text/html; charset=euc-kr">',
    '<meta charset="euc-kr">',
    "",                                                           # no declaration: guessed
])
def test_html_legacy_korean(head):
    src = f"<!DOCTYPE html><html><head>{head}<title>제목</title></head><body><p>안녕하세요 똠방각하</p></body></html>"
    out = translate_html(src.encode("cp949"), _prefix).decode("utf-8")
    assert "<title>[KO] 제목</title>" in out and "<p>[KO] 안녕하세요 똠방각하</p>" in out
    assert out.count("charset") == 1 and '<meta charset="utf-8">' in out and "euc-kr" not in out


def test_html_chinese_without_meta():
    src = "<p>这是一个中文文本文件，用于测试编码检测。</p>".encode("gbk")
    assert _collect(translate_html, src) == ["这是一个中文文本文件，用于测试编码检测。"]


def test_html_fallback_keeps_atoms():
    src = b"<p>Run <code>make all</code> then see <img src=a.png> the result.</p>"
    # <x1> instead of <x1/>: recovered in place
    out = translate_html(src, lambda xs: ["실행 <x1> then see <x2> the result." for _ in xs]).decode()
    assert '<p>실행 <code>make all</code> then see <img src="a.png"> the result.</p>' in out
    # unusable tags: plain text, but the code sample and the image are kept
    out = translate_html(src, lambda xs: ["실행 <x3> 그리고 <g9>결과</g9>." for _ in xs]).decode()
    assert "<code>make all</code>" in out and '<img src="a.png">' in out and "실행" in out


# ====================================================================== tag fallback entities
def test_tag_fallback_unescapes_entities():
    xml = (f'<w:p xmlns:w="{ooxml.W}"><w:r><w:t xml:space="preserve">R&amp;D budget </w:t></w:r>'
           '<w:r><w:footnoteReference w:id="1"/></w:r><w:r><w:t xml:space="preserve"> rose</w:t></w:r></w:p>')
    p = etree.fromstring(xml)
    seen = []

    def tr(xs):
        seen.extend(xs)
        return ["R&amp;D 예산이 증가했다" for _ in xs]            # every tag dropped
    stats = translate_paragraphs(DocxAdapter(), [p], tr)
    assert seen == ["R&amp;D budget <x1/> rose"] and stats["fallback"] == 1
    assert "".join(t.text or "" for t in p.iter(w("t"))) == "R&D 예산이 증가했다"
    assert p.find(f".//{w('footnoteReference')}") is not None
    assert preview_from_pairs([(seen[0], "R&amp;D 예산이 증가했다")]) == "R&D 예산이 증가했다"
    items = build_report([(seen[0], "R&amp;D 예산이 증가했다")], source_lang="en", target_lang="ko")
    assert items and all(i["target"] == "R&D 예산이 증가했다" for i in items)


# ====================================================================== report
@pytest.mark.parametrize(("src", "tgt"), [
    ("KRW 35,000,000", "3,500만 원"),
    ("120,000,000,000 won", "1,200억 원"),
    ("2.3 trillion won", "2.3조 원"),
    ("2,300,000,000,000 won", "2.3조 원"),
    ("2024년 3월 15일", "March 15, 2024"),
    ("2024.03.15 기준", "as of March 15, 2024"),
    ("2024. 3. 15. 기준", "as of March 15, 2024"),
    ("2024.03 기준", "as of March 2024"),
])
def test_report_no_false_number_alarms(src, tgt):
    assert missing_numbers(src, tgt) == []


def test_report_still_flags_wrong_numbers():
    assert missing_numbers("KRW 35,000,000", "3,600만 원") == ["35000000"]
    assert missing_numbers("2024년 3월 16일", "March 15, 2024") == ["16"]
    assert missing_numbers("3개월 동안", "for 4 months") == ["3"]


# ====================================================================== Markdown / subtitles
def test_markdown_leading_rule_is_not_front_matter():
    assert _collect(translate_markdown, "---\n\nIntro paragraph after a rule.\n\nSecond paragraph.\n") == [
        "Intro paragraph after a rule.", "Second paragraph."]
    assert _collect(translate_markdown, "---\n\nIntro.\n\n---\n\nMore text.\n") == ["Intro.", "More text."]
    assert _collect(translate_markdown, "---\ntitle: test\ntags:\n  - a\n---\n# Safety\n") == ["Safety"]


def test_markdown_trailing_hash_only_closes_headings():
    src = "# Using C#\nWe write services in C#\n## Closing ##  \nIssue #\n"
    assert _collect(translate_markdown, src) == ["Using C#", "We write services in C#", "Closing", "Issue #"]
    out = translate_markdown(src, _prefix)
    assert out == "# [KO] Using C#\n[KO] We write services in C#\n## [KO] Closing ##  \n[KO] Issue #\n"


def test_subtitles_whitespace_only_separator():
    src = ("1\n00:00:01,000 --> 00:00:03,000\nHello there\n \n"
           "2\n00:00:04,000 --> 00:00:06,000\nGood morning\n\t\n\n3\n00:00:07,000 --> 00:00:08,000\nBye\n")
    assert _collect(translate_subtitles, src) == ["Hello there", "Good morning", "Bye"]
    out = translate_subtitles(src, _prefix)
    assert "2\n00:00:04,000 --> 00:00:06,000\n[KO] Good morning\n" in out


# ====================================================================== DOCX bilingual copy
def test_docx_clone_drops_style_numbering_and_page_break():
    xml = (f'<w:p xmlns:w="{ooxml.W}"><w:pPr><w:pStyle w:val="Heading1"/><w:pageBreakBefore/>'
           '<w:spacing w:after="0"/><w:jc w:val="center"/></w:pPr><w:r><w:t>Introduction</w:t></w:r></w:p>')
    c = docx_clone_paragraph(etree.fromstring(xml))
    ppr = c.find(w("pPr"))
    assert [etree.QName(ch).localname for ch in ppr] == ["pStyle", "pageBreakBefore", "numPr", "spacing", "jc"]
    assert ppr.find(w("pageBreakBefore")).get(w("val")) == "0"
    assert ppr.find(f"{w('numPr')}/{w('numId')}").get(w("val")) == "0"
    # paragraph without pPr: still gets the overrides (style numbering may come from the default style)
    c = docx_clone_paragraph(etree.fromstring(f'<w:p xmlns:w="{ooxml.W}"><w:r><w:t>Body</w:t></w:r></w:p>'))
    assert [etree.QName(ch).localname for ch in c.find(w("pPr"))] == ["pageBreakBefore", "numPr"]


# ====================================================================== HWPX bilingual copy
def _outline_hwpx() -> bytes:
    src = F.make_hwpx()
    hp = F.HP_NS
    para_pr = ('<hh:paraPr id="{id}"><hh:heading type="{h}" idRef="0" level="0"/>'
               '<hh:breakSetting keepWithNext="0" pageBreakBefore="0"/></hh:paraPr>')
    header = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes" ?>'
        f'<hh:head xmlns:hh="{HH}" version="1.4" secCnt="1"><hh:refList>'
        '<hh:charProperties itemCnt="1"><hh:charPr id="0" height="1000"/></hh:charProperties>'
        '<hh:paraProperties itemCnt="2">' + para_pr.format(id=0, h="NONE") + para_pr.format(id=1, h="OUTLINE")
        + '</hh:paraProperties></hh:refList></hh:head>')
    section = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes" ?>'
        f'<hs:sec xmlns:hp="{hp}" xmlns:hs="http://www.hancom.co.kr/hwpml/2011/section">'
        '<hp:p id="1" paraPrIDRef="0" styleIDRef="0" pageBreak="0" columnBreak="0" merged="0">'
        '<hp:run charPrIDRef="0"><hp:ctrl><hp:header id="1" applyPageType="BOTH"><hp:subList id="">'
        '<hp:p id="0" paraPrIDRef="0" styleIDRef="0" pageBreak="0" columnBreak="0" merged="0">'
        '<hp:run charPrIDRef="0"><hp:t>Hydrogen report header</hp:t></hp:run></hp:p>'
        '</hp:subList></hp:header></hp:ctrl></hp:run>'
        '<hp:run charPrIDRef="0"><hp:t>Overview of the program</hp:t></hp:run></hp:p>'
        '<hp:p id="2" paraPrIDRef="1" styleIDRef="2" pageBreak="1" columnBreak="0" merged="0">'
        '<hp:run charPrIDRef="0"><hp:t>Background and purpose</hp:t></hp:run></hp:p></hs:sec>')
    buf = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(src)) as zin, zipfile.ZipFile(buf, "w") as zout:
        for info in zin.infolist():
            data = {"Contents/header.xml": header.encode(), "Contents/section0.xml": section.encode()}.get(
                info.filename, zin.read(info.filename))
            zout.writestr(info, data)
    return buf.getvalue()


def test_hwpx_bilingual_copy_has_no_numbering_or_page_break(tmp_path):
    src, dst = tmp_path / "a.hwpx", tmp_path / "b.hwpx"
    src.write_bytes(_outline_hwpx())
    assert set(_collect(translate_hwpx, src, None, bilingual=True)) == {
        "Overview of the program", "Background and purpose", "Hydrogen report header"}
    translate_hwpx(src, dst, _prefix, bilingual=True)
    out = dst.read_bytes()
    with zipfile.ZipFile(io.BytesIO(out)) as z:
        assert z.infolist()[0].filename == "mimetype"
        head = etree.fromstring(z.read("Contents/header.xml"))
        sec = etree.fromstring(z.read("Contents/section0.xml"))
    props = head.find(f".//{{{HH}}}paraProperties")
    twins = {p.get("id"): p for p in props}
    assert props.get("itemCnt") == "3" and set(twins) == {"0", "1", "2"}
    assert twins["2"].find(f"{{{HH}}}heading").get("type") == "NONE"
    assert twins["1"].find(f"{{{HH}}}heading").get("type") == "OUTLINE"        # original untouched
    hp = "{%s}" % F.HP_NS
    body = sec.findall(f"{hp}p")
    texts = ["".join(p.itertext()) for p in body]
    i = texts.index("Background and purpose")
    copy_ = body[i + 1]
    assert texts[i + 1] == "[KO] Background and purpose"
    assert body[i].get("pageBreak") == "1" and copy_.get("pageBreak") == "0"
    assert body[i].get("paraPrIDRef") == "1" and copy_.get("paraPrIDRef") == "2"
    # the running header is translated in place, not left in the source language or duplicated
    hdr = [t.text for t in sec.iter(f"{hp}t") if t.text and "header" in t.text]
    assert hdr == ["[KO] Hydrogen report header"]


# ====================================================================== XLSX
def _xlsx_with_table() -> bytes:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.worksheet.table import Table

    wb = Workbook()
    ws = wb.active
    ws.append(["Region", "Value", "Note"])
    ws.append(["East", 10, "Strong growth"])
    ws.append(["West", 20, "Stable demand"])
    ws.add_table(Table(displayName="Table1", ref="A1:C3"))
    ws["E1"] = "=SUM(Table1[Value])"
    ch = BarChart()
    ch.title = "Value by region"
    ch.add_data(Reference(ws, min_col=2, min_row=1, max_row=3), titles_from_data=True)
    ws.add_chart(ch, "G2")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_xlsx_table_headers_kept_and_chart_title_translated(tmp_path):
    from openpyxl import load_workbook

    src, dst = tmp_path / "t.xlsx", tmp_path / "t_ko.xlsx"
    src.write_bytes(_xlsx_with_table())
    translate_xlsx(src, dst, _prefix)
    ws = load_workbook(dst).active
    assert [c.value for c in ws[1][:3]] == ["Region", "Value", "Note"]          # = tableColumn names
    assert ws["A2"].value == "[KO] East" and ws["C3"].value == "[KO] Stable demand"
    table = F.zip_text(dst.read_bytes(), "xl/tables/table1.xml")
    assert 'name="Region"' in table and 'name="Value"' in table
    assert "[KO] Value by region" in F.zip_text(dst.read_bytes(), "xl/charts/chart1.xml")


def test_xlsx_threaded_comments(tmp_path):
    tc = ooxml.TC
    xml = (f'<ThreadedComments xmlns="{tc}"><threadedComment ref="A1" id="{{1}}"><text>Please check</text>'
           f'</threadedComment><threadedComment ref="A1" id="{{2}}" parentId="{{1}}"><text>@Kim done</text>'
           '<mentions><mention startIndex="0" length="4"/></mentions></threadedComment></ThreadedComments>')
    src = io.BytesIO(F.make_minimal_xlsx())
    buf = io.BytesIO()
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(buf, "w") as zout:
        for info in zin.infolist():
            zout.writestr(info, zin.read(info.filename))
        zout.writestr("xl/threadedComments/threadedComment1.xml", xml)
    (tmp_path / "c.xlsx").write_bytes(buf.getvalue())
    translate_xlsx(tmp_path / "c.xlsx", tmp_path / "c_ko.xlsx", _prefix)
    out = F.zip_text((tmp_path / "c_ko.xlsx").read_bytes(), "xl/threadedComments/threadedComment1.xml")
    assert "<text>[KO] Please check</text>" in out
    assert "<text>@Kim done</text>" in out                       # mention offsets stay valid


def test_xlsx_sheets_without_inline_strings_are_not_parsed(tmp_path, monkeypatch):
    src = tmp_path / "m.xlsx"
    src.write_bytes(F.make_minimal_xlsx())
    parsed: list[str] = []
    real = ooxml.rewrite_zip

    def spy(s, d, select, transform):
        def tf(name, tree):
            parsed.append(name)
            transform(name, tree)
        return real(s, d, select, tf)
    monkeypatch.setattr(ooxml, "rewrite_zip", spy)
    translate_xlsx(src, tmp_path / "m_ko.xlsx", _prefix)
    assert "xl/worksheets/sheet1.xml" not in parsed and "xl/sharedStrings.xml" in parsed
    with zipfile.ZipFile(src) as a, zipfile.ZipFile(tmp_path / "m_ko.xlsx") as b:
        assert a.read("xl/worksheets/sheet1.xml") == b.read("xl/worksheets/sheet1.xml")


@pytest.mark.parametrize("chunk", [1, 3, 7, 64, 1 << 20])
def test_xlsx_formula_scan_across_chunk_boundaries(tmp_path, monkeypatch, chunk):
    monkeypatch.setattr(ooxml, "_SCAN_CHUNK", chunk)
    src = tmp_path / "p.xlsx"
    src.write_bytes(F.make_xlsx())
    lits, inline = ooxml.scan_sheets(src)
    assert {"Yes", "Running", "Stopped", "No"} <= lits
    with zipfile.ZipFile(src) as z:
        sheet = z.read("xl/worksheets/sheet1.xml")
    assert inline == ({"xl/worksheets/sheet1.xml"} if b"inlineStr" in sheet else set())


def test_xlsx_formula_scan_entities_and_prefixes(tmp_path):
    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    sheet = (f'<x:worksheet xmlns:x="{ns}"><x:sheetData><x:row r="1"><x:c r="A1" t="inlineStr"><x:is><x:t>Done'
             '</x:t></x:is></x:c><x:c r="B1"><x:f t="shared" si="0"/></x:c>'
             '<x:c r="C1"><x:f>IF(A1=&quot;Done&quot;,"R&amp;D","x")</x:f></x:c></x:row></x:sheetData>'
             '<x:conditionalFormatting><x:cfRule><x:formula>$A1="Open"</x:formula></x:cfRule>'
             '</x:conditionalFormatting></x:worksheet>')
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("xl/worksheets/sheet1.xml", sheet)
    (tmp_path / "s.xlsx").write_bytes(buf.getvalue())
    lits, inline = ooxml.scan_sheets(tmp_path / "s.xlsx")
    assert lits == {"Done", "R&D", "x", "Open"} and inline == {"xl/worksheets/sheet1.xml"}


# ====================================================================== charts (PPTX / DOCX)
def test_pptx_chart_titles_translated(tmp_path):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Hydrogen output"
    cd = CategoryChartData()
    cd.categories = ["East", "West"]
    cd.add_series("Output", (10, 20))
    ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.5), Inches(6), Inches(4), cd).chart
    ch.has_title = True
    ch.chart_title.text_frame.text = "Output by region"
    prs.save(tmp_path / "c.pptx")
    translate_pptx(tmp_path / "c.pptx", tmp_path / "c_ko.pptx", _prefix, set_lang="ko-KR")
    out = Presentation(tmp_path / "c_ko.pptx")
    chart = next(sh for sh in out.slides[0].shapes if sh.has_chart).chart
    assert chart.chart_title.text_frame.text == "[KO] Output by region"
